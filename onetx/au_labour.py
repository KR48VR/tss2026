"""Australia: the third workforce. Employment and earnings by ANZSCO unit group, the ISCO-08 link to
O*NET, firm AI use from the ABS business survey, the industry mix for sector shifts, and the ABS's
own task lists per unit group.

Sources (data/raw/abs/, see docs/DATA_SOURCES.md):
  lfs_EQ08.xlsx                 Labour Force, Australia, Detailed: employed persons by ANZSCO unit group
                                (4-digit, ANZSCO 2013 v1.2), sex and state, quarterly since 1986
  lfs_EQ09.xlsx                 the same survey: employed persons by ANZSIC industry division and ANZSCO
                                major group
  eeh_63060DO011_2025/2305.xlsx Employee Earnings and Hours, May 2025 and May 2023, data cube 11: average
                                weekly total cash earnings of all employees by ANZSCO unit group
  cab_BCSDC06_2024-25.xlsx      Characteristics of Australian Business 2024-25, cube 6 table 2: share of
                                businesses using each technology, including AI, by industry and size
  cab_81670DO007_202122.xlsx    the same survey for 2021-22 (cube 7 table 1)
  anzsco2022_structure.xlsx     ANZSCO 2022: groups, titles and skill levels
  anzsco2022_correspondence.xlsx table 3: ISCO-08 unit group -> ANZSCO 2022 occupations (6-digit)
  anzsco_pages/<unit>.html      the ABS's classification pages, one per unit group, with the official
                                "Tasks Include" list (fetch.fetch_anzsco_pages)

ANZSCO 2013 v1.2 (the labour force cubes) and ANZSCO 2022 (earnings, correspondence) agree on all but
18 unit-group codes; a 2022 code the labour force cube does not carry is mapped to the 2013 codes that
share its minor group (first three digits).
"""
from __future__ import annotations

import datetime as dt
import html as htmllib
import re
from pathlib import Path

import numpy as np
import pandas as pd

from . import io
from .config import RAW

ABS = RAW / "abs"
PAGES = ABS / "anzsco_pages"

MAJOR_GROUP_NAMES = {"1": "Managers", "2": "Professionals", "3": "Technicians and trades workers", "4": "Community and personal service workers",
                     "5": "Clerical and administrative workers", "6": "Sales workers", "7": "Machinery operators and drivers", "8": "Labourers"}
# ANZSIC 2006 divisions, as the labour force and business surveys label them
DIVISIONS = {"A": "Agriculture, Forestry and Fishing", "B": "Mining", "C": "Manufacturing", "D": "Electricity, Gas, Water and Waste Services", "E": "Construction",
             "F": "Wholesale Trade", "G": "Retail Trade", "H": "Accommodation and Food Services", "I": "Transport, Postal and Warehousing",
             "J": "Information Media and Telecommunications", "K": "Financial and Insurance Services", "L": "Rental, Hiring and Real Estate Services",
             "M": "Professional, Scientific and Technical Services", "N": "Administrative and Support Services", "O": "Public Administration and Safety",
             "P": "Education and Training", "Q": "Health Care and Social Assistance", "R": "Arts and Recreation Services", "S": "Other Services"}
SKILL_LEVEL_NAMES = {1: "bachelor degree or higher, or five years' experience", 2: "diploma or advanced diploma, or three years' experience", 3: "certificate IV or III with two years' training, or three years' experience",
                     4: "certificate II or III, or one year's experience", 5: "certificate I or secondary schooling, with short on-the-job training"}


def _rows(path: Path, sheet: str, min_row: int = 1) -> list[tuple]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return list(wb[sheet].iter_rows(min_row=min_row, values_only=True))


def _num(v):
    try:
        x = float(str(v).replace(",", "").strip())
        return x if np.isfinite(x) else np.nan
    except (TypeError, ValueError):
        return np.nan


def _division(label: str) -> str | None:
    s = str(label).strip().lower()
    for k, v in DIVISIONS.items():
        if s.startswith(v.lower()[:20]):
            return k
    return None


# ----------------------------------------------------------------------------- classification
def load_anzsco_structure() -> pd.DataFrame:
    """One row per ANZSCO 2022 unit group: code, title, major/sub-major/minor groups and names, skill level
    (the mean of the levels listed when a unit group spans two)."""
    rows = _rows(ABS / "anzsco2022_structure.xlsx", "Table 4", 6)
    names = {}
    out = []
    for r in rows:
        vals = [x for x in r if x is not None and str(x).strip() != ""]
        if not vals:
            continue
        code = str(vals[0]).strip()
        if not code.isdigit():
            continue
        title = str(vals[1]).strip() if len(vals) > 1 else ""
        names[code] = title
        if len(code) == 4:
            lv = str(vals[2]) if len(vals) > 2 else ""
            levels = [int(x) for x in re.findall(r"\d", lv)]
            out.append({"anzsco_code": code, "anzsco_title": title, "major_group": code[0], "submajor": code[:2], "minor": code[:3], "skill_level": float(np.mean(levels)) if levels else np.nan, "skill_levels_listed": lv.strip()})
    df = pd.DataFrame(out)
    df["submajor_name"] = df.submajor.map(names)
    df["minor_name"] = df.minor.map(names)
    df["major_group_name"] = df.major_group.map(MAJOR_GROUP_NAMES)
    return df


def load_isco_to_anzsco(unit_codes: set[str] | None = None) -> pd.DataFrame:
    """ISCO-08 unit group -> ANZSCO unit group (4-digit) pairs from the ABS correspondence table (6-digit
    occupations rolled up). Weight = 1/n over the distinct unit groups an ISCO code maps to. When
    `unit_codes` (the codes the labour force cube carries) is given, a 2022 unit group absent from it is
    replaced by the carried codes of its minor group."""
    rows = _rows(ABS / "anzsco2022_correspondence.xlsx", "Table 3", 7)
    pairs, cur, cur_title = [], None, None
    for r in rows:
        if r[0] is not None and re.fullmatch(r"\d{4}", str(r[0]).strip()):
            cur, cur_title = str(r[0]).strip(), str(r[1]).strip() if r[1] else ""
        code6 = str(r[2]).strip() if r[2] is not None else ""
        if cur and re.fullmatch(r"\d{6}", code6):
            pairs.append({"isco_code": cur, "isco_title": cur_title, "anzsco_code": code6[:4], "partial": str(r[3]).strip().lower() == "p" if r[3] else False})
    df = pd.DataFrame(pairs)
    if unit_codes:
        unit_codes = set(unit_codes)
        fixed = []
        for _, r in df.iterrows():
            if r.anzsco_code in unit_codes:
                fixed.append(r.to_dict())
            else:
                sibs = sorted(c for c in unit_codes if c[:3] == r.anzsco_code[:3] and not c.endswith("0"))
                for c in sibs:
                    fixed.append({**r.to_dict(), "anzsco_code": c, "remapped_from_2022": r.anzsco_code})
        df = pd.DataFrame(fixed)
    g = df.groupby(["isco_code", "anzsco_code"]).agg(isco_title=("isco_title", "first"), partial=("partial", "max")).reset_index()
    g["weight"] = 1.0 / g.groupby("isco_code").anzsco_code.transform("count")
    return g


# ----------------------------------------------------------------------------- employment and earnings
def load_lfs_occupation_quarterly() -> pd.DataFrame:
    """Employed persons (thousands) by ANZSCO unit group and quarter, both sexes and all states summed.
    Parsing the 800,000-row cube takes a minute, so the result is cached as a processed table."""
    if io.exists("au_lfs_occupation_quarterly"):
        return io.load("au_lfs_occupation_quarterly")
    import openpyxl

    wb = openpyxl.load_workbook(ABS / "lfs_EQ08.xlsx", read_only=True, data_only=True)
    agg: dict[tuple, float] = {}
    labels = {}
    for row in wb["Data 1"].iter_rows(min_row=5, values_only=True):
        if not isinstance(row[0], dt.datetime) or row[1] not in ("Males", "Females"):
            continue
        label = str(row[3])
        code = label[:4]
        if not code.isdigit():
            continue
        labels[code] = label[5:].strip()
        key = (row[0].date(), code)
        agg[key] = agg.get(key, 0.0) + float(row[4] or 0.0)
    df = pd.DataFrame([{"quarter": pd.Timestamp(q), "anzsco_code": c, "label": labels[c], "employed_thousands": v} for (q, c), v in agg.items()]).sort_values(["quarter", "anzsco_code"]).reset_index(drop=True)
    io.save(df, "au_lfs_occupation_quarterly")
    return df


def redistribute_nfd(emp: pd.Series) -> pd.Series:
    """'Not further defined' codes (ending in 0: 1000, 1100, 1110) hold people the survey could not place in
    a unit group. Their employment is shared out over the unit groups under the same prefix, in proportion
    to those groups' own employment, from the finest level up."""
    s = emp.copy().astype(float)
    nfd = sorted([c for c in s.index if c.endswith("0")], key=lambda c: -len(c.rstrip("0")))
    for c in nfd:
        prefix = c.rstrip("0")
        kids = [k for k in s.index if not k.endswith("0") and k.startswith(prefix)]
        tot = float(s[kids].sum()) if kids else 0.0
        if kids and tot > 0:
            s[kids] = s[kids] + s[c] * s[kids] / tot
        elif kids:
            s[kids] = s[kids] + s[c] / len(kids)
        s = s.drop(c)
    return s


def annual_employment(year: int, quarterly: pd.DataFrame | None = None) -> pd.Series:
    """Mean of the year's quarters per unit group, thousands, nfd codes redistributed."""
    q = load_lfs_occupation_quarterly() if quarterly is None else quarterly
    d = q[q.quarter.dt.year == year]
    emp = d.groupby("anzsco_code").employed_thousands.sum() / d.quarter.nunique()
    return redistribute_nfd(emp)


def load_eeh_earnings(year: int = 2025) -> pd.DataFrame:
    """Average weekly total cash earnings (A$) of all employees by ANZSCO unit group, persons."""
    path = ABS / f"eeh_63060DO011_{year}05.xlsx"
    out = []
    for r in _rows(path, "Table_1", 7):
        if r[0] is None:
            continue
        m = re.match(r"(\d{4})\s+(.*)", str(r[0]).strip())
        if not m:
            continue
        out.append({"anzsco_code": m.group(1), "eeh_title": m.group(2).strip(), "weekly_earnings": _num(r[3]), "year": year})
    return pd.DataFrame(out)


def build_anzsco_weights(year: int = 2025) -> pd.DataFrame:
    """Per unit group (labour force codes): employment (mean of the year's quarters, thousands), weekly
    earnings (May 2025 survey; minor, sub-major or major group mean where the group is too small to
    publish), their product as the wage weight, skill level and group names."""
    emp = annual_employment(year)
    q = load_lfs_occupation_quarterly()
    labels = q.drop_duplicates("anzsco_code").set_index("anzsco_code").label
    struct = load_anzsco_structure().set_index("anzsco_code")
    eeh = load_eeh_earnings(2025).dropna(subset=["weekly_earnings"]).set_index("anzsco_code").weekly_earnings
    df = pd.DataFrame({"anzsco_code": emp.index, "employment_thousands": emp.values})
    df["anzsco_title"] = df.anzsco_code.map(labels)
    df["major_group"] = df.anzsco_code.str[0]
    df["submajor_label"] = df.anzsco_code.str[:2]
    df["minor"] = df.anzsco_code.str[:3]
    sub_names = struct.drop_duplicates("submajor").set_index("submajor").submajor_name
    df["submajor_name"] = df.submajor_label.map(sub_names)
    df["skill_level"] = df.anzsco_code.map(struct.skill_level)
    # a 2013 code missing from the 2022 structure takes its minor group's typical level
    for lvl, n in [("minor", 3), ("submajor_label", 2), ("major_group", 1)]:
        grp = struct.groupby(struct.index.str[:n]).skill_level.mean()
        df["skill_level"] = df.skill_level.fillna(df[lvl].map(grp))
    df["wage_weekly"] = df.anzsco_code.map(eeh)
    df["wage_source"] = np.where(df.wage_weekly.notna(), "eeh_2025_unit_group", None)
    for lvl, n in [("minor", 3), ("submajor_label", 2), ("major_group", 1)]:
        grp = eeh.groupby(eeh.index.str[:n]).mean()
        m = df.wage_weekly.isna()
        df.loc[m, "wage_weekly"] = df.loc[m, lvl].map(grp)
        df.loc[m & df.wage_weekly.notna(), "wage_source"] = f"eeh_2025_{lvl.replace('_label', '')}_group_mean"
    m = df.wage_weekly.isna()
    df.loc[m, "wage_weekly"] = float(eeh.mean())
    df.loc[m, "wage_source"] = "eeh_2025_all_mean"
    df["wage_value"] = df.employment_thousands * df.wage_weekly
    df["employment_method"] = "lfs_unit_group_annual_mean_nfd_redistributed"
    df["edition"] = year
    return df.drop(columns=["minor"]).reset_index(drop=True)


def employment_change(year_from: int = 2024, year_to: int = 2025) -> pd.DataFrame:
    """Rung five for Australia: employment change between two years' quarterly means (labour force
    survey) and the change in average weekly earnings between the May 2023 and May 2025 earnings
    surveys, per unit group."""
    e0, e1 = annual_employment(year_from), annual_employment(year_to)
    w0 = load_eeh_earnings(2023).dropna(subset=["weekly_earnings"]).set_index("anzsco_code").weekly_earnings
    w1 = load_eeh_earnings(2025).dropna(subset=["weekly_earnings"]).set_index("anzsco_code").weekly_earnings
    df = pd.DataFrame({f"emp_{year_from}": e0, f"emp_{year_to}": e1}).reset_index().rename(columns={"index": "anzsco_code"})
    df["emp_change_pct"] = 100.0 * (df[f"emp_{year_to}"] / df[f"emp_{year_from}"].where(df[f"emp_{year_from}"] > 0) - 1.0)
    df["wage_2023"] = df.anzsco_code.map(w0)
    df["wage_2025"] = df.anzsco_code.map(w1)
    df["wage_change_pct"] = 100.0 * (df.wage_2025 / df.wage_2023.where(df.wage_2023 > 0) - 1.0)
    df["year_from"], df["year_to"] = year_from, year_to
    df["wage_years"] = "2023-2025"
    return df


# ----------------------------------------------------------------------------- industry mix and firm AI use
def load_industry_by_major(year: int = 2025) -> pd.DataFrame:
    """Employed persons (thousands, full- plus part-time) by ANZSIC division and ANZSCO major group, the
    mean of the year's quarters."""
    name_to_digit = {v.lower(): k for k, v in MAJOR_GROUP_NAMES.items()}
    agg: dict[tuple, float] = {}
    quarters = set()
    for row in _rows(ABS / "lfs_EQ09.xlsx", "Data 1", 5):
        if not isinstance(row[0], dt.datetime) or row[0].year != year or row[1] not in ("Males", "Females"):
            continue
        div = _division(row[2])
        mg = name_to_digit.get(str(row[3]).strip().lower())
        if div is None or mg is None:
            continue
        quarters.add(row[0].date())
        agg[(div, mg)] = agg.get((div, mg), 0.0) + float(row[4] or 0.0) + float(row[5] or 0.0)
    n = max(1, len(quarters))
    return pd.DataFrame([{"industry": d, "industry_name": DIVISIONS[d], "major_group": mg, "employed_thousands": v / n, "year": year} for (d, mg), v in agg.items()])


def _cab_table(path: Path, sheet: str, header_row: int, ai_header: str, period: str, date: float) -> pd.DataFrame:
    rows = _rows(path, sheet, 1)
    hdr = rows[header_row - 1]
    cells = [str(c).strip().lower() if c is not None else "" for c in hdr]
    ai_col = next(i for i, c in enumerate(cells) if c.startswith(ai_header))
    out, section = [], None
    for r in rows[header_row:]:
        label = str(r[0]).strip() if r[0] is not None else ""
        if not label:
            continue
        if label in ("Employment size", "Industry", "State/Territory", "Location", "Footnotes"):
            section = label
            continue
        if label == "Total":
            out.append({"period": period, "date": date, "group_type": "all", "industry": "ALL", "label": "All businesses", "pct": _num(r[ai_col])})
        elif section == "Industry":
            div = _division(label)
            if div:
                out.append({"period": period, "date": date, "group_type": "industry", "industry": div, "label": DIVISIONS[div], "pct": _num(r[ai_col])})
        elif section == "Employment size":
            out.append({"period": period, "date": date, "group_type": "size", "industry": label, "label": label, "pct": _num(r[ai_col])})
    return pd.DataFrame(out)


def load_cab_ai_use() -> pd.DataFrame:
    """Share of businesses that used AI during the financial year, nationally, by ANZSIC division and by
    employment size: 2021-22 (cube 7) and 2024-25 (cube 6). The ABS did not publish 2022-23 or 2023-24
    editions. A year that ends in June is dated at its end (mid-year), the conservative choice against
    the US question (use in the last two weeks)."""
    frames = []
    p1 = ABS / "cab_81670DO007_202122.xlsx"
    if p1.exists():
        frames.append(_cab_table(p1, "Table_1", 6, "artificial intelligence", "2021-22", 2022.5))
    p2 = ABS / "cab_BCSDC06_2024-25.xlsx"
    if p2.exists():
        frames.append(_cab_table(p2, "Table 2", 5, "artificial intelligence", "2024-25", 2025.5))
    if not frames:
        raise FileNotFoundError("ABS Characteristics of Australian Business cubes (data/raw/abs/cab_*.xlsx)")
    return pd.concat(frames, ignore_index=True)


# ----------------------------------------------------------------------------- the ABS's own task lists
def load_anzsco_pages() -> pd.DataFrame:
    """Per unit group from the ABS classification pages: the lead description, the indicative skill level
    text, and the official 'Tasks Include' items (one row per task)."""
    if not PAGES.exists():
        raise FileNotFoundError(f"{PAGES} (run fetch.fetch_anzsco_pages)")
    out = []
    for p in sorted(PAGES.glob("*.html")):
        code = p.stem
        if not re.fullmatch(r"\d{4}", code):
            continue
        t = p.read_text(encoding="utf-8", errors="ignore")
        m = re.search(r"Tasks Include:</h3>.*?<ul>(.*?)</ul>", t, re.S)
        tasks = []
        if m:
            for li in re.findall(r"<li>(.*?)</li>", m.group(1), re.S):
                txt = htmllib.unescape(re.sub(r"<[^>]+>", " ", li))
                txt = re.sub(r"\s+", " ", txt).strip().rstrip(".")
                if len(txt) > 3:
                    tasks.append(txt)
        desc = ""
        md = re.search(r"<h1[^>]*>.*?</h1>.*?<p>(.*?)</p>", t, re.S)
        if md:
            desc = re.sub(r"\s+", " ", htmllib.unescape(re.sub(r"<[^>]+>", " ", md.group(1)))).strip()
        sk = re.search(r"Indicative Skill Level:.*?<p>(.*?)</p>", t, re.S)
        skill_text = re.sub(r"\s+", " ", htmllib.unescape(re.sub(r"<[^>]+>", " ", sk.group(1)))).strip() if sk else ""
        for i, task in enumerate(tasks):
            out.append({"anzsco_code": code, "task_seq": i, "task": task, "description": desc, "skill_text": skill_text})
        if not tasks:
            out.append({"anzsco_code": code, "task_seq": -1, "task": None, "description": desc, "skill_text": skill_text})
    return pd.DataFrame(out)


def load_anzsco_tasks() -> pd.DataFrame:
    """One row per official task item: anzsco_code, task_seq, task (unit groups whose page lists none are left out)."""
    d = load_anzsco_pages()
    return d[d.task.notna()][["anzsco_code", "task_seq", "task"]].reset_index(drop=True)


def load_title_index() -> pd.DataFrame:
    """The ABS index of principal titles, alternative titles and specialisations (6-digit occupations),
    rolled up to the unit group: columns unit, title, category."""
    out = []
    for r in _rows(ABS / "anzsco2022_index.xlsx", "Table 2", 7):
        if r[0] is None or not re.fullmatch(r"\d{6}", str(r[0]).strip()):
            continue
        out.append({"unit": str(r[0]).strip()[:4], "title": str(r[1]).strip(), "category": str(r[2]).strip() if r[2] else ""})
    return pd.DataFrame(out)
